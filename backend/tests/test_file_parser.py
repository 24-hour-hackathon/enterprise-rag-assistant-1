import os
import tempfile
import fitz  # PyMuPDF
import docx
from app.utils.file_parser import parse_file


def test_parse_txt_file():
    with tempfile.NamedTemporaryFile(suffix=".txt", delete=False, mode="w", encoding="utf-8") as f:
        f.write("# Enterprise Leave Policy\n\nEmployees are entitled to 20 days of paid annual leave.\n\n# Sick Leave\n\nEmployees get 10 sick leaves per year.")
        temp_path = f.name

    try:
        doc = parse_file(temp_path)
        assert doc.file_type == "txt"
        assert len(doc.blocks) >= 2
        assert "Enterprise Leave Policy" in doc.raw_text
        assert any(b.section == "Enterprise Leave Policy" or b.section == "Sick Leave" for b in doc.blocks)
    finally:
        os.remove(temp_path)


def test_parse_pdf_file():
    # Create a small valid PDF in memory
    pdf_doc = fitz.open()
    page1 = pdf_doc.new_page()
    page1.insert_text((50, 50), "Company Overview\nAcme Corp is an enterprise technology provider.")
    page2 = pdf_doc.new_page()
    page2.insert_text((50, 50), "Code of Conduct\nIntegrity and security are paramount.")

    with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as f:
        temp_path = f.name

    try:
        pdf_doc.save(temp_path)
        pdf_doc.close()

        doc = parse_file(temp_path)
        assert doc.file_type == "pdf"
        assert doc.total_pages == 2
        assert len(doc.blocks) >= 2
        assert doc.blocks[0].page_number == 1
        assert doc.blocks[-1].page_number == 2
    finally:
        if os.path.exists(temp_path):
            os.remove(temp_path)


def test_parse_docx_file():
    docx_doc = docx.Document()
    docx_doc.add_heading("Health Insurance Benefits", level=1)
    docx_doc.add_paragraph("Comprehensive medical and dental coverage is provided to all full-time personnel.")

    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as f:
        temp_path = f.name

    try:
        docx_doc.save(temp_path)
        doc = parse_file(temp_path)
        assert doc.file_type == "docx"
        assert len(doc.blocks) >= 1
        assert "Health Insurance Benefits" in doc.raw_text
    finally:
        if os.path.exists(temp_path):
            os.remove(temp_path)
